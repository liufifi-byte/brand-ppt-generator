"""mdpptx 单元测试 — 不调 LLM (解析/规划/折行/检查/渲染后处理/模板渲染)。"""
from pathlib import Path
import sys

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from mdpptx.md_parser import parse_md
from mdpptx.planner import plan_pages
from mdpptx.template_mode import _wrap, render_cover, render_news, render_end
from mdpptx.checks import (check_xml, check_clean, check_decoration,
                           check_verbatim, _nums)
from mdpptx.render import post_fix
from mdpptx.theme import get_theme

SAMPLE_MD = """__第3期新能源产业市场周报（10/06-10/10）__

__市场价格动态__

10月6日，电池级碳酸锂现货价格为8.65万元/吨，本周内基本保持稳定，10月10日，现货价格收于8.70万元/吨。

10月月度长期合同均价为9.20万元/吨，较9月上涨0.15万元/吨。

__行业政策与市场动态__

- __三部门联合印发新型储能制造业发展方案__（2026/10/08）

近日，工业和信息化部等三部门联合印发实施方案，提出到2028年产值规模突破1.2万亿元。

方案支持多技术路线并行发展。

- __欧洲议会通过动力电池碳足迹核算新规__（2026/10/07）

欧洲议会以412票赞成通过修正案。

新规将使出口欧洲的电池产品平均成本上升4.5%左右。
"""


@pytest.fixture
def doc():
    return parse_md(SAMPLE_MD)


# ── md 解析 ──
def test_parse_md_structure(doc):
    assert doc["doc_title"] == "第3期新能源产业市场周报（10/06-10/10）"
    names = [m["name"] for m in doc["modules"]]
    assert names == ["市场价格动态", "行业政策与市场动态"]
    price_mod = doc["modules"][0]
    assert len(price_mod["paras"]) == 2 and not price_mod.get("items")
    news_mod = doc["modules"][1]
    assert len(news_mod["items"]) == 2
    it = news_mod["items"][0]
    assert it["title"] == "三部门联合印发新型储能制造业发展方案"
    assert it["date"] == "2026/10/08"
    assert len(it["paras"]) == 2


def test_parse_md_unescapes():
    d = parse_md("__标题\\-转义__\n\n__模块__\n\n- __条目\\(一)\\[__（2026/10/08）\n\n正文。")
    assert d["doc_title"] == "标题-转义"
    assert d["modules"][0]["items"][0]["title"] == "条目(一)["


# ── 页面规划 ──
def test_plan_pages(doc):
    pages = plan_pages(doc)
    types = [p["type"] for p in pages]
    assert types == ["cover", "price", "news", "news", "end"]


# ── 折行: 数字串原子性 ──
def test_wrap_never_splits_numbers():
    text = "现货价格从8.65万元/吨涨至103,250元区间附近"
    lines = _wrap(text, 200)   # 故意窄宽迫使折行
    joined = "".join(lines)
    assert joined == text
    # 每行中不得出现被切断的数字 (重组后完整)
    assert "8.65" in joined and "103,250" in joined


def test_wrap_cjk_width():
    # "一二三四五六七八九十"×2, fs=20 → 每行 5 个, 共 4 行
    assert _wrap("一二三四五六七八九十一二三四五六七八九十", 100) == \
        ["一二三四五", "六七八九十", "一二三四五", "六七八九十"]


# ── 检查器 ──
def test_check_xml():
    assert check_xml("<svg><rect/></svg>")[0]
    assert not check_xml("<svg><rect/>")[0]


def test_check_clean_patterns():
    ok, issues = check_clean('<text>source: xx</text>')
    assert not ok and any("source" in i for i in issues)
    ok, _ = check_clean('<text fill="#004A86">正常</text>')
    assert ok


def test_check_decoration_min_imgs():
    svg4 = '<svg>' + '<image href="x"/>' * 4 + '</svg>'
    svg2 = '<svg>' + '<image href="x"/>' * 2 + '</svg>'
    assert check_decoration(svg4, 4)[0]
    assert not check_decoration(svg2, 4)[0]
    assert check_decoration(svg2, 2)[0]


def test_check_verbatim_numbers():
    src = ["产量达到1234吨，增长5.6%。"]
    good = '<svg><text>产量达到1234吨，增长5.6%。</text></svg>'
    bad_num = '<svg><text>产量达到1234吨，增长7.8%。</text></svg>'
    assert check_verbatim(good, src)[0]
    ok, issues = check_verbatim(bad_num, src)
    assert not ok and any("素材外" in i for i in issues)


def test_check_verbatim_material_missing():
    src = ["这一段很长很长很长应该逐字出现"]
    svg = '<svg><text>这一段被改写了很长很长</text></svg>'
    ok, issues = check_verbatim(svg, src)
    assert not ok and any("逐字" in i for i in issues)


def test_nums_decimal_tokens():
    assert _nums("刻度 8.4 与 9.0") == {"8.4", "9.0"}


# ── 渲染后处理 ──
def test_post_fix_strips_fence_and_injects_decor():
    T = get_theme()
    raw = "```xml\n<svg xmlns=\"http://www.w3.org/2000/svg\">\n<text>hi</text>\n</svg>\n```"
    svg = post_fix(raw)
    assert svg.startswith("<?xml")
    assert 'id="decoration"' in svg
    assert "```" not in svg
    assert T["decor"].split("\n")[1] in svg   # image1 装饰在


def test_post_fix_strips_fullpage_bg_rect_and_animate():
    raw = ('<svg xmlns="http://www.w3.org/2000/svg">'
           '<g id="decoration"><image/></g>'
           '<rect x="0" y="0" width="1280" height="720" fill="#fff"/>'
           '<animate attributeName="x" to="10"/></svg>')
    svg = post_fix(raw)
    assert 'width="1280" height="720"' not in svg
    assert "<animate" not in svg


# ── 模板渲染 ──
def test_render_cover_news_end_valid_xml():
    total = 5
    cover = render_cover("第3期新能源产业市场周报", "2026年10月6日至10月10日", 1, total)
    news = render_news("行业政策与市场动态", "三部门联合印发方案", "2026/10/08",
                       ["第一段正文内容较长需要折行处理一下。", "第二段正文。"],
                       [{"v": "1.2万亿", "label": "产值目标"}], "补充信息一句。", 2, total)
    end = render_end(total, total)
    for svg in (cover, news, end):
        assert check_xml(svg)[0], svg[:200]
    assert "1.2万亿" in news and "补充：" in news
    assert "5 / 5" in end
    # 封面用封面装饰 (image5/image6), 内容页用标准装饰 (image1-4)
    assert "image5" in cover and "image1" in news


def test_render_news_overflow_protection():
    long_para = "很长的正文段落。" * 60
    svg = render_news("模块", "标题", "2026/10/08",
                      [long_para, "第二段" + "字" * 200], [], "", 2, 8)
    assert check_xml(svg)[0]
    assert "2 / 8" in svg


# ── 文档接入: pdf/docx 文本提取 ──
def _make_synthetic_pdf(path):
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 90), "Synthetic Market Report", fontsize=16, fontname="helv")
    page.insert_text((72, 130), "The price reached 8.65 units in October.", fontsize=11, fontname="helv")
    doc.save(str(path)); doc.close()


def _make_synthetic_docx(path):
    from docx import Document
    d = Document()
    d.add_paragraph("Synthetic Market Report")
    d.add_paragraph("The price reached 8.65 units in October.")
    d.save(str(path))


def test_extract_text_pdf_and_docx(tmp_path):
    from mdpptx.ingest import extract_text
    pdf = tmp_path / "a.pdf"; _make_synthetic_pdf(pdf)
    docx = tmp_path / "a.docx"; _make_synthetic_docx(docx)
    t1 = extract_text(pdf)
    assert "8.65" in t1 and "Synthetic" in t1
    t2 = extract_text(docx)
    assert "8.65" in t2 and "Synthetic" in t2


def test_structured_ok_rejects_plain_text():
    from mdpptx.ingest import _structured_ok
    assert not _structured_ok("只是普通文本没有结构")
    assert _structured_ok("__标题__\n\n__模块__\n\n正文段落足够长可以判断。")


# ── PPT 模版主题色提取 ──
def test_extract_theme_override_from_pptx(tmp_path):
    from mdpptx.template_pptx import extract_theme_override
    from pptx import Presentation
    p = tmp_path / "t.pptx"
    Presentation().save(str(p))       # 默认模板自带 clrScheme
    ov = extract_theme_override(p)
    assert "colors" in ov and "primary" in ov["colors"]
    assert ov["colors"]["primary"].startswith("#") and len(ov["colors"]["primary"]) == 7


def test_theme_override_apply_restore_cycle():
    from mdpptx import theme
    from mdpptx.template_mode import C
    snap = theme.snapshot()
    theme.apply_override({"colors": {"primary": "#ABCDEF"}, "name": "测试"})
    assert theme.THEME["colors"]["primary"] == "#ABCDEF" and C["primary"] == "#ABCDEF"
    theme.restore(snap)
    assert C["primary"] == "#004A86" and theme.THEME["name"] != "测试"
